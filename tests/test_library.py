import base64
import io
from datetime import UTC, datetime

import pytest
from fastapi.testclient import TestClient

from kb import db, library, scopes
from kb.api import auth
from kb.api.app import build_app
from kb.routines import factpack

CSRF = {"X-KB-Request": "1"}
NOW = datetime(2026, 10, 20, 9, 0, tzinfo=UTC)


def tiny_pdf(text: str) -> bytes:
    """A one-page PDF with real text (Helvetica), small enough to build by hand."""
    content = f"BT /F1 12 Tf 20 100 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 400 144] /Contents 4 0 R"
        b" /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = io.BytesIO(), []
    out.write(b"%PDF-1.4\n")
    for n, body in enumerate(objects, 1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % n + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for off in offsets:
        out.write(b"%010d 00000 n \n" % off)
    out.write(
        b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    )
    return out.getvalue()


def tiny_docx() -> bytes:
    import docx

    d = docx.Document()
    d.add_heading("Principles", level=1)
    d.add_paragraph("I take new clients only through referrals.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_extract_md_pdf_docx():
    assert (
        library.extract_text("plan.md", b"# Goals\nLaunch Epicentr") == "# Goals\nLaunch Epicentr"
    )
    assert "Constitution principle" in library.extract_text(
        "c.pdf", tiny_pdf("Constitution principle")
    )
    text = library.extract_text("c.docx", tiny_docx())
    assert "# Principles" in text and "only through referrals" in text
    with pytest.raises(library.LibraryError, match="unsupported"):
        library.extract_text("photo.jpg", b"...")


def test_document_requires_mark(conn):
    with pytest.raises(library.LibraryError, match="choose personal or a project"):
        library.add_document(conn, "Plan", "plan", [], "p.md", b"x", scopes.known(conn), NOW)
    with pytest.raises(library.LibraryError, match="unknown scope"):
        library.add_document(conn, "Plan", "plan", ["acme"], "p.md", b"x", scopes.known(conn), NOW)
    with pytest.raises(library.LibraryError, match="kind"):
        library.add_document(
            conn, "Plan", "diary", ["personal"], "p.md", b"x", scopes.known(conn), NOW
        )
    assert conn.execute("SELECT count(*) AS n FROM documents").fetchone()["n"] == 0


def test_versions_latest_indexed(conn):
    doc = library.add_document(
        conn,
        "Constitution",
        "constitution",
        ["personal"],
        "c.md",
        b"alpha rule",
        scopes.known(conn),
        NOW,
    )
    library.add_version(conn, doc, "c.md", b"beta rule", NOW)
    item = conn.execute(
        "SELECT raw_text, kind, metadata FROM items WHERE source_item_key = %s", (f"doc-{doc}",)
    ).fetchone()
    assert item["raw_text"] == "beta rule" and item["kind"] == "document"
    assert item["metadata"]["private"] is False  # core documents: AI may read by default
    versions = conn.execute(
        "SELECT text FROM document_versions WHERE document_id = %s ORDER BY id", (doc,)
    ).fetchall()
    assert [v["text"] for v in versions] == ["alpha rule", "beta rule"]


def test_long_document_sections_point_to_the_document(conn):
    body = (
        "\n".join(f"# Section {i}\n" + ("words " * 80) for i in range(4))
        + "\n# Clients\nOnly referrals from dealers."
    )
    doc = library.add_document(
        conn,
        "Constitution",
        "constitution",
        ["personal"],
        "c.md",
        body.encode(),
        scopes.known(conn),
        NOW,
    )
    parent = conn.execute(
        "SELECT id FROM items WHERE source_item_key = %s", (f"doc-{doc}",)
    ).fetchone()["id"]
    sections = conn.execute(
        "SELECT title, metadata FROM items WHERE kind = 'doc_section' ORDER BY id"
    ).fetchall()
    assert len(sections) == 5 and sections[-1]["title"] == "Constitution - Clients"
    assert all(s["metadata"]["parent_item_id"] == parent for s in sections)
    # a new version replaces the sections
    library.add_version(conn, doc, "c.md", b"short now", NOW)
    assert (
        conn.execute("SELECT count(*) AS n FROM items WHERE kind = 'doc_section'").fetchone()["n"]
        == 0
    )


def test_documents_stay_out_of_briefs(conn):
    library.add_document(
        conn, "Plan", "plan", ["personal"], "p.md", b"plan text", scopes.known(conn), NOW
    )
    pack = factpack.build(
        conn, "morning", (datetime(2026, 1, 1, tzinfo=UTC), datetime(2027, 1, 1, tzinfo=UTC))
    )
    assert pack.by_scope == {}


@pytest.fixture
def client(conn, pg_url, config):
    app = build_app(
        None,
        connect=lambda: db.connect(pg_url),
        rp=auth.RelyingParty("https://os.x"),
        config=config,
    )
    c = TestClient(app, base_url="https://os.x")
    c.cookies.set(auth.SESSION_COOKIE, auth.create_session(conn, "t", datetime.now(UTC)))
    return c


def test_library_api_upload_note_search_and_versions(client, conn):
    pdf = base64.b64encode(tiny_pdf("Quokka plan launch Epicentr")).decode()
    r = client.post(
        "/api/library",
        json={
            "kind": "plan",
            "scopes": ["quokka"],
            "title": "Quokka plan 2026",
            "file": {"filename": "q.pdf", "content_b64": pdf},
        },
        headers=CSRF,
    )
    assert r.status_code == 200
    doc = r.json()["id"]
    note = client.post(
        "/api/library",
        json={"kind": "note", "scopes": ["personal"], "title": "Health", "text": "checkup notes"},
        headers=CSRF,
    )
    assert note.status_code == 200
    docs = client.get("/api/library").json()["documents"]
    assert [d["title"] for d in docs] == ["Quokka plan 2026", "Health"]  # core documents first
    assert docs[1]["ai_allowed"] is False  # notes are private by default
    assert (
        client.post(
            "/api/library",
            json={"kind": "note", "scopes": [], "text": "x", "title": "x"},
            headers=CSRF,
        ).status_code
        == 400
    )
    hits = client.get("/api/search", params={"q": "Epicentr"}).json()["results"]
    assert [h["title"] for h in hits] == ["Quokka plan 2026"]
    md = base64.b64encode(b"Launch Rozetka first").decode()
    client.post(
        f"/api/library/{doc}/versions",
        json={"file": {"filename": "q.md", "content_b64": md}},
        headers=CSRF,
    )
    detail = client.get(f"/api/library/{doc}").json()
    assert detail["text"] == "Launch Rozetka first" and len(detail["versions"]) == 2
    old = detail["versions"][1]["id"]
    assert "Epicentr" in client.get(f"/api/library/{doc}/versions/{old}").json()["text"]
    assert (
        client.patch(f"/api/library/{doc}", json={"ai_allowed": False}, headers=CSRF).status_code
        == 200
    )
    meta = conn.execute(
        "SELECT metadata FROM items WHERE source_item_key = %s", (f"doc-{doc}",)
    ).fetchone()["metadata"]
    assert meta["private"] is True
    assert (
        client.post(
            "/api/library",
            json={
                "kind": "plan",
                "scopes": ["personal"],
                "file": {"filename": "a.md", "content_b64": "%%%"},
            },
            headers=CSRF,
        ).status_code
        == 400
    )
